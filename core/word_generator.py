import os
import copy
from typing import Dict, Any, Optional
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

def gerar_documento_word(
    dados_viagem: Dict[str, Any],
    template_path: str,
    output_path: str,
    linhas_em_branco_extras: int = 3,
    tamanho_fonte_tabela: int = 12
) -> str:
    """
    Gera o documento Word preenchido com base no template e dados extraídos.
    - Mantém a numeração padrão nativa do Word (1., 2., 3...) na coluna N sem duplicar texto.
    - Ajusta a altura das linhas limpando parágrafos extras.
    - Aplica o tamanho de fonte 12 na tabela.
    - Preserva o cabeçalho e rodapé oficiais.
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Arquivo modelo não encontrado: {template_path}")

    doc = docx.Document(template_path)

    destino = dados_viagem.get("destino", "MONTES CLAROS")
    data_viagem = dados_viagem.get("data_viagem", "")
    dia_semana = dados_viagem.get("dia_semana", "")
    hora_saida = dados_viagem.get("hora_saida", "05:00")
    passageiros = dados_viagem.get("passageiros", [])

    # 1. Atualiza os parágrafos de cabeçalho no corpo do documento
    if len(doc.paragraphs) > 1:
        doc.paragraphs[1].text = f"VIAGEM  DE {destino} {data_viagem}  {dia_semana}"
        for r in doc.paragraphs[1].runs:
            r.bold = True
            r.font.size = Pt(16)
            
    if len(doc.paragraphs) > 2:
        doc.paragraphs[2].text = f"SAÍDA - {hora_saida} HRS DO POSTO DE SAUDE"
        for r in doc.paragraphs[2].runs:
            r.bold = True
            r.font.size = Pt(16)

    # 2. Preenche a tabela
    if not doc.tables:
        raise ValueError("O documento modelo não contém nenhuma tabela.")

    tbl = doc.tables[0]
    
    # Prepara o molde da linha mantendo formatação original
    if len(tbl.rows) > 1:
        template_tr = copy.deepcopy(tbl.rows[1]._tr)
    else:
        template_tr = copy.deepcopy(tbl.rows[0]._tr)

    # Remove altura fixa forçada para que a altura se ajuste naturalmente ao conteúdo
    trPr = template_tr.xpath(".//w:trPr")
    if trPr:
        for h in trPr[0].xpath(".//w:trHeight"):
            trPr[0].remove(h)

    # Limpa parágrafos extras de cada célula do molde para evitar linhas vazias
    for tc in template_tr.xpath(".//w:tc"):
        p_elements = tc.xpath(".//w:p")
        # Mantém apenas o primeiro parágrafo
        for p_elem in p_elements[1:]:
            tc.remove(p_elem)
        # Limpa o texto do primeiro parágrafo
        first_p = tc.xpath(".//w:p")[0]
        for r in first_p.xpath(".//w:r"):
            first_p.remove(r)

    # Remove todas as linhas de dados antigas, mantendo apenas o cabeçalho (linha 0)
    for r in tbl.rows[1:]:
        tbl._tbl.remove(r._tr)

    # Função auxiliar para configurar a célula com apenas 1 parágrafo limpo
    def set_cell(cell, text: str, align=WD_ALIGN_PARAGRAPH.LEFT, bold: bool = False):
        # Garante que não haja parágrafos extras que causam espaçamento vertical indesejado
        while len(cell.paragraphs) > 1:
            p_elem = cell.paragraphs[-1]._p
            cell._tc.remove(p_elem)

        if not cell.paragraphs:
            cell.add_paragraph()
        
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0

        # Remove runs existentes
        for r in p.runs:
            p._p.remove(r._r)

        if text:
            run = p.add_run(text)
            run.bold = bold
            run.font.size = Pt(tamanho_fonte_tabela)

    # Preenche cada passageiro
    for p_data in passageiros:
        new_tr = copy.deepcopy(template_tr)
        tbl._tbl.append(new_tr)
        row = tbl.rows[-1]

        # Coluna 0: N (Deixa o texto vazio para usar exclusivamente a numeração nativa "1.", "2." do Word)
        set_cell(row.cells[0], "", align=WD_ALIGN_PARAGRAPH.CENTER)

        # Coluna 1: HRS (Hora do Atendimento)
        set_cell(row.cells[1], p_data.get("hora", ""), align=WD_ALIGN_PARAGRAPH.CENTER)

        # Coluna 2: PACIENTE
        tipo = p_data.get("tipo", "").lower()
        nome = p_data.get("nome", "").strip()
        if tipo == "acompanhante":
            if not nome or nome.upper() == "ACOMPANHANTE" or nome == "-":
                pac_text = "ACOMP."
            else:
                pac_text = f"ACOMP. {nome}"
        else:
            pac_text = nome
        set_cell(row.cells[2], pac_text, align=WD_ALIGN_PARAGRAPH.LEFT)

        # Coluna 3: CONTATO (Em branco ou 'VOLTA' se passageiro somente de volta)
        if p_data.get("is_volta_only", False):
            set_cell(row.cells[3], "VOLTA", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        else:
            set_cell(row.cells[3], "", align=WD_ALIGN_PARAGRAPH.CENTER)

        # Coluna 4: CLINICA (Estabelecimento)
        set_cell(row.cells[4], p_data.get("estabelecimento", ""), align=WD_ALIGN_PARAGRAPH.LEFT)

        # Coluna 5: PROCEDIMENTO (Em branco)
        set_cell(row.cells[5], "", align=WD_ALIGN_PARAGRAPH.CENTER)

        # Coluna 6: C (Em branco)
        set_cell(row.cells[6], "", align=WD_ALIGN_PARAGRAPH.CENTER)

        # Coluna 7: LOCAL (Em branco)
        set_cell(row.cells[7], "", align=WD_ALIGN_PARAGRAPH.LEFT)

    # 3. Adiciona linhas extras em branco no final para anotações manuais se necessário
    for _ in range(linhas_em_branco_extras):
        new_tr = copy.deepcopy(template_tr)
        tbl._tbl.append(new_tr)
        row = tbl.rows[-1]
        for cell in row.cells:
            set_cell(cell, "")

    # Garante que o diretório de destino exista
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    doc.save(output_path)
    return output_path
